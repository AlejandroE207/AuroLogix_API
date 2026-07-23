import csv
import io
import json
import re
import unicodedata
from collections import OrderedDict
from decimal import Decimal, InvalidOperation

from app.repository import picking_repository, order_repository, inventory_repository, motion_repository
from app.model.picking_model import Orden, Picking, Picking_detalle, Picking_items
from app.model.inventory_model import Motion
from datetime import datetime


REQUIRED_IMPORT_COLUMNS = {
    "codigo": {"nro documento", "nro. documento", "numero documento", "codigo", "código"},
    "cliente": {"razon social cliente factura", "cliente", "razon social", "razón social cliente factura"},
    "cod_item": {"referencia"},
    "cantidad": {"cantidad", "cant. pendiente", "cant pendiente", "cantidad pendiente"},
    "bodega": {"bodega"},
}

OPTIONAL_IMPORT_COLUMNS = {
    "fecha": {"fecha"},
    "estado": {"estado movto.", "estado movto", "estado"},
    "lote": {"lote"},
    "detalles": {"cliente factura", "detalles", "detalle"},
}


def _normalise_header(value) -> str:
    value = unicodedata.normalize("NFD", str(value or "").strip().casefold())
    value = "".join(character for character in value if unicodedata.category(character) != "Mn")
    return re.sub(r"\s+", " ", value)


def _normalise_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_blank_row(row: dict) -> bool:
    """Distingue una fila vacía de una fila de pedido incompleta."""
    return all(not _normalise_value(value) for value in row.values())


def _parse_quantity(value) -> int:
    raw = _normalise_value(value).replace(" ", "")
    if not raw:
        raise ValueError("la cantidad está vacía")
    if "," in raw and "." in raw:
        raw = raw.replace(".", "").replace(",", ".")
    elif "," in raw:
        raw = raw.replace(",", ".")
    quantity = Decimal(raw)
    if quantity <= 0 or quantity != quantity.to_integral_value():
        raise ValueError("la cantidad debe ser un entero mayor que cero")
    return int(quantity)


def _allocate_inventory(requested_items: list[dict], available: dict[int, list[dict]]):
    """Asigna primero el lote solicitado y luego lotes alternos por FEFO."""
    tasks, errors = [], []
    for requested_item in requested_items:
        item_id = requested_item["id_item"]
        requested_lot = _normalise_value(requested_item.get("lote"))
        pending = int(requested_item["cantidad"])

        entries = available.get(item_id, [])
        ordered_entries = sorted(
            entries,
            key=lambda entry: (
                0 if requested_lot and _normalise_value(entry.get("lote")).casefold() == requested_lot.casefold() else 1,
                entry.get("fecha_vencimiento") is None,
                str(entry.get("fecha_vencimiento") or "9999-12-31"),
            ),
        )
        for entry in ordered_entries:
            if pending == 0:
                break
            assigned = min(pending, int(entry["cantidad"]))
            if assigned <= 0:
                continue
            tasks.append({
                "id_item": item_id,
                "cantidad": assigned,
                "lote": entry["lote"],
                "id_posicion_origen": entry["id_posicion"],
            })
            entry["cantidad"] -= assigned
            pending -= assigned

        if pending:
            errors.append({
                "id_item": item_id,
                "item_code": requested_item.get("item_code"),
                "lote": requested_lot,
                "faltante": pending,
            })
    return tasks, errors


def _read_import_rows(content: bytes, extension: str) -> list[dict]:
    if extension == "csv":
        try:
            text_content = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text_content = content.decode("latin-1")
        try:
            dialect = csv.Sniffer().sniff(text_content[:4096], delimiters=";,\t")
        except csv.Error:
            return list(csv.DictReader(io.StringIO(text_content), delimiter=";"))
        return list(csv.DictReader(io.StringIO(text_content), dialect=dialect))
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            return []
        return [dict(zip(headers, row)) for row in rows if any(value is not None for value in row)]
    except Exception as error:
        raise ValueError(f"No fue posible leer el archivo XLSX: {error}") from error


def _column_mapping(rows: list[dict]) -> dict[str, str]:
    if not rows:
        raise ValueError("El archivo no contiene filas de datos")
    headers = {_normalise_header(header): header for header in rows[0].keys()}
    mapping, missing = {}, []
    for field, aliases in REQUIRED_IMPORT_COLUMNS.items():
        header = next(
            (headers[_normalise_header(alias)] for alias in aliases if _normalise_header(alias) in headers),
            None,
        )
        if header is None:
            missing.append(field)
        else:
            mapping[field] = header
    if missing:
        labels = {
            "codigo": "Nro documento",
            "cliente": "Razón social cliente factura",
            "cod_item": "Referencia",
            "cantidad": "Cantidad",
            "bodega": "Bodega",
        }
        raise ValueError("Faltan columnas obligatorias: " + ", ".join(labels[item] for item in missing))
    for field, aliases in OPTIONAL_IMPORT_COLUMNS.items():
        mapping[field] = next(
            (headers[_normalise_header(alias)] for alias in aliases if _normalise_header(alias) in headers),
            None,
        )
    return mapping


async def import_orders_file(db, content: bytes, extension: str, id_usuario: int) -> dict:
    """Importa pedidos aprobados, agrupándolos por documento de venta."""
    try:
        rows = _read_import_rows(content, extension)
        columns = _column_mapping(rows)
    except ValueError as error:
        return {"result": 0, "message": str(error)}

    orders: OrderedDict[str, dict] = OrderedDict()
    errors, skipped_rows = [], 0
    for row_number, row in enumerate(rows, start=2):
        # Los CSV/XLSX exportados desde Excel suelen conservar filas vacías al
        # final. No representan pedidos ni deben impedir la importación.
        if _is_blank_row(row):
            skipped_rows += 1
            continue
        bodega = _normalise_value(row.get(columns["bodega"])) if columns["bodega"] else ""
        # Solo se importan las filas de bodegas cuyo nombre o código comienza por M.
        if not bodega.casefold().startswith("m"):
            skipped_rows += 1
            continue
        status = _normalise_value(row.get(columns["estado"])) if columns["estado"] else ""
        if status and _normalise_header(status) != "aprobada":
            skipped_rows += 1
            continue
        codigo = _normalise_value(row.get(columns["codigo"]))
        cliente = _normalise_value(row.get(columns["cliente"]))
        cod_item = _normalise_value(row.get(columns["cod_item"]))
        lote = _normalise_value(row.get(columns["lote"])) if columns["lote"] else ""
        detalles = _normalise_value(row.get(columns["detalles"])) if columns["detalles"] else ""
        try:
            cantidad = _parse_quantity(row.get(columns["cantidad"]))
        except (InvalidOperation, ValueError) as error:
            errors.append(f"Fila {row_number}: {error}")
            continue
        if not codigo or not cliente or not cod_item:
            errors.append(f"Fila {row_number}: documento, cliente e Id item son obligatorios")
            continue
        order = orders.setdefault(codigo, {
            "codigo": codigo, "cliente": cliente, "items": OrderedDict(),
            "fecha": _normalise_value(row.get(columns["fecha"])) if columns["fecha"] else "",
            "bodega": bodega,
            "detalles": detalles,
        })
        if order["cliente"] != cliente:
            errors.append(f"Fila {row_number}: el documento {codigo} tiene clientes distintos")
            continue
        item_key = (cod_item, lote)
        order["items"][item_key] = order["items"].get(item_key, 0) + cantidad
    if not orders:
        return {"result": 0, "message": "No hay filas aprobadas para importar"}
    if errors:
        return {"result": 0, "message": "El archivo tiene errores y no se importó ninguna orden: " + "; ".join(errors[:10])}

    existing_codes = await picking_repository.get_existing_order_codes(db, list(orders))
    not_registered_orders = [
        {"codigo": code, "motivo": "La orden ya existe"}
        for code in orders
        if code in existing_codes
    ]
    new_orders = [order for code, order in orders.items() if code not in existing_codes]

    item_codes = list({code for order in new_orders for code, _ in order["items"]})
    items_by_code = await picking_repository.get_items_by_codes(db, item_codes)
    orders_to_create = []
    for order in new_orders:
        missing_codes = sorted({code for code, _ in order["items"] if code not in items_by_code})
        if missing_codes:
            not_registered_orders.append({
                "codigo": order["codigo"],
                "motivo": "No existen artículos para los códigos: " + ", ".join(missing_codes),
            })
            continue

        requested_items = []
        for (item_code, requested_lot), requested in order["items"].items():
            requested_items.append({
                "id_item": items_by_code[item_code]["id"],
                "cantidad": requested,
                "lote": requested_lot,
            })
        details = {
            "fecha": order["fecha"],
            "bodega": order["bodega"],
            "detalles": order["detalles"],
            "origen": "importacion_archivo",
        }
        orders_to_create.append({"codigo": order["codigo"], "cliente": order["cliente"], "id_usuario": id_usuario,
                                 "tipo": "OF", "detalles": json.dumps(details, ensure_ascii=False), "items": requested_items})

    if not orders_to_create:
        return {
            "result": 1,
            "message": "El archivo fue procesado; no había órdenes nuevas para registrar.",
            "created_orders": 0,
            "created_pickings": 0,
            "skipped_rows": skipped_rows,
            "registered_orders": [],
            "not_registered_orders": not_registered_orders,
        }
    try:
        created = await picking_repository.create_bulk_pickings(db, orders_to_create)
        await db.commit()
    except Exception as error:
        await db.rollback()
        return {"result": 0, "message": f"No fue posible crear las órdenes: {error}"}
    registered_orders = [order["codigo"] for order in created]
    return {
        "result": 1,
        "message": f"Se registraron {len(created)} órdenes de salida.",
        "created_orders": len(created),
        "created_pickings": len(created),
        "skipped_rows": skipped_rows,
        "registered_orders": registered_orders,
        "not_registered_orders": not_registered_orders,
        "orders": created,
    }

async def create_picking(db, orden: Orden, picking_data: Picking, picking_detalle: Picking_detalle):
    requested_items = []
    for item in picking_detalle.items:
        if item.id_item is None or item.cantidad is None or item.cantidad <= 0:
            return Picking(
                result=0,
                message="Cada item debe tener id_item y una cantidad mayor que cero",
            )
        requested_items.append({
            "id_item": item.id_item,
            "cantidad": item.cantidad,
            "lote": item.lote,
        })

    inventory = await picking_repository.get_inventory_for_items(
        db, list({item["id_item"] for item in requested_items})
    )
    available = {
        item_id: [dict(entry) for entry in entries]
        for item_id, entries in inventory.items()
    }
    allocated_tasks, allocation_errors = _allocate_inventory(requested_items, available)
    if allocation_errors:
        descriptions = [
            f"item {error['id_item']}"
            + (f", lote {error['lote']}" if error["lote"] else "")
            + f": faltan {error['faltante']} unidades"
            for error in allocation_errors
        ]
        return Picking(
            result=0,
            message="Inventario insuficiente: " + "; ".join(descriptions),
        )

    create_order = await order_repository.create_order(db, orden)
    if create_order.result == 1:
        picking_data.id_orden = create_order.id
        create_picking_data = await picking_repository.create_picking(db, picking_data)
        if create_picking_data.result == 1:
            picking_detalle.id_picking = create_picking_data.id
            create_picking_detalle = await picking_repository.create_picking_detalle(db, picking_detalle)
            if create_picking_detalle.result == 1:
                picking_tasks = Picking_detalle(
                    id_picking=create_picking_data.id,
                    items=[Picking_items(**task) for task in allocated_tasks],
                )
                create_picking_task = await picking_repository.create_picking_task(db, picking_tasks)
                if create_picking_task.result == 1:
                    return create_picking_data
                else:
                    return create_picking_data
            else:
                return create_picking_data
        else:
            return create_picking_data
    else:
        return create_order
        # return picking_data
    


async def get_next_task(db, id_picking: int):
    """Obtiene la siguiente tarea pendiente del picking."""
    task = await picking_repository.get_next_task(db, id_picking)
    if task.result == 0:
        task_count = await picking_repository.count_picking_tasks(db, id_picking)
        if task_count == 0:
            details = await picking_repository.get_picking_details(db, id_picking)
            if not details:
                return {
                    "completed": False,
                    "message": "El picking no tiene productos asociados.",
                }

            requested_items = [
                {
                    "id_item": detail["id_item"],
                    "item_code": str(detail["cod_item"]),
                    "cantidad": detail["cantidad"],
                    "lote": detail["lote"],
                }
                for detail in details
            ]
            inventory = await picking_repository.get_inventory_for_items(
                db, list({item["id_item"] for item in requested_items})
            )
            available = {
                item_id: [dict(entry) for entry in entries]
                for item_id, entries in inventory.items()
            }
            tasks, allocation_errors = _allocate_inventory(requested_items, available)
            if allocation_errors:
                missing_inventory = [
                    {
                        "cod_item": error["item_code"],
                        "lote": error["lote"] or None,
                        "cantidad_faltante": error["faltante"],
                    }
                    for error in allocation_errors
                ]
                return {
                    "completed": False,
                    "awaiting_stock": True,
                    "message": "El picking está pendiente por inventario insuficiente.",
                    "missing_inventory": missing_inventory,
                }

            try:
                await picking_repository.insert_picking_tasks(db, id_picking, tasks)
                await db.commit()
            except Exception as error:
                await db.rollback()
                return {
                    "completed": False,
                    "message": f"No fue posible generar las tareas de picking: {error}",
                }

            task = await picking_repository.get_next_task(db, id_picking)
            if task.result == 1:
                return {"completed": False, "task": task}
            return {
                "completed": False,
                "message": "No fue posible obtener la tarea de picking generada.",
            }

        result = await picking_repository.complete_picking(db, id_picking)
        if result["result"] == 1:
            order_data = await order_repository.update_order_status_by_picking(db, id_picking)
            if order_data.result == 1:
                return {"completed":True, "message": "No hay más tareas pendientes. Picking completado."}
            else:
                return {"completed":False, "message": "Error al actualizar el estado de la orden"}
        else:
            return {"completed":False, "message": "Error al completar el picking"}
    return {"completed":False , "task": task}

async def get_tasks_overview(db, id_picking: int):
    """Devuelve todas las extracciones y cualquier cantidad aún sin inventario."""
    details = await picking_repository.get_picking_details(db, id_picking)
    if not details:
        return {"result": 0, "message": "El picking no tiene productos asociados"}

    task_count = await picking_repository.count_picking_tasks(db, id_picking)
    missing_inventory = []
    if task_count == 0:
        requested_items = [
            {
                "id_item": detail["id_item"],
                "item_code": str(detail["cod_item"]),
                "cantidad": detail["cantidad"],
                "lote": detail["lote"],
            }
            for detail in details
        ]
        inventory = await picking_repository.get_inventory_for_items(
            db, list({item["id_item"] for item in requested_items})
        )
        available = {
            item_id: [dict(entry) for entry in entries]
            for item_id, entries in inventory.items()
        }
        tasks, errors = _allocate_inventory(requested_items, available)
        if tasks:
            await picking_repository.insert_picking_tasks(db, id_picking, tasks)
            await db.commit()
        missing_inventory = [
            {
                "id_item": error["id_item"],
                "cod_item": error["item_code"],
                "lote_solicitado": error["lote"] or None,
                "cantidad_faltante": error["faltante"],
            }
            for error in errors
        ]

    tasks = await picking_repository.list_picking_tasks(db, id_picking)
    requested_by_item = {}
    for detail in details:
        entry = requested_by_item.setdefault(
            detail["id_item"],
            {"cod_item": str(detail["cod_item"]), "cantidad": 0, "lotes": []},
        )
        entry["cantidad"] += int(detail["cantidad"])
        lot = _normalise_value(detail.get("lote"))
        if lot and lot not in entry["lotes"]:
            entry["lotes"].append(lot)

    assigned_by_item = {}
    for task in tasks:
        assigned_by_item[task["id_item"]] = (
            assigned_by_item.get(task["id_item"], 0) + int(task["cantidad"])
        )
    missing_inventory = [
        {
            "id_item": item_id,
            "cod_item": requested["cod_item"],
            "lotes_solicitados": requested["lotes"],
            "cantidad_faltante": requested["cantidad"] - assigned_by_item.get(item_id, 0),
        }
        for item_id, requested in requested_by_item.items()
        if requested["cantidad"] > assigned_by_item.get(item_id, 0)
    ]
    return {
        "result": 1,
        "message": "Extracciones del picking obtenidas exitosamente",
        "tasks": tasks,
        "missing_inventory": missing_inventory,
        "has_missing_inventory": bool(missing_inventory),
    }


# ACTUALIZAR ESTADO DE ORDEN ACA MISMO, DESPUES DE CONFIRMAR LA PRIMERA TAREA, SE CAMBIA EL ESTADO DE LA ORDEN
# Y CUANDO SE FINALICE Y NO HAYA UNA TAREA MAS SE FINALIZA LA ORDEN
async def confirm_task(
    db,
    id_task: int,
    id_picking: int,
    cod_posicion_escaneada: str,
    id_usuario: int,
):
    """
    Confirma una tarea de picking después de escanear la posición de origen.
    Registra el movimiento y actualiza el inventario.
    """
    # Obtener la información de la tarea
    task = await picking_repository.get_task_by_id(db, id_picking, id_task)
    
    if task is None:
        return {"result": 0, "message": "Tarea no encontrada"}

    if task.estado != 1:
        return {
            "result": 0,
            "message": "La tarea seleccionada ya fue realizada o no está pendiente",
        }
    
    # Validar que la posición escaneada sea correcta
    if cod_posicion_escaneada != task.cod_posicion_origen:
        return {"result": 0, "message": f"Posición incorrecta. Se esperaba {task.id_posicion_origen}, se escaneó {cod_posicion_escaneada}"}
    
    try:
        # Crear registro de movimiento (salida/picking)
        # tipo_movimiento: 1=ENTRADA, 2=SALIDA
        cantidad_salida = float(task.cantidad)
        motion = Motion(
            tipo_movimiento=2,  # SALIDA
            id_item=task.id_item,
            lote=task.lote,
            cantidad=cantidad_salida,
            posicion_origen_id=task.id_posicion_origen,
            posicion_destino_id=None,  # No hay posición destino en picking
            fecha=datetime.now(),
            id_usuario=id_usuario
        )
        
        motion_result = await motion_repository.create_motion_output(
            db, motion, commit=False
        )
        
        if motion_result.result == 0:
            raise RuntimeError("Error al registrar el movimiento")
        
        # Actualizar inventario (restar la cantidad)
        inventory_result = await inventory_repository.update_inventory_quantity(
            db,
            task.id_posicion_origen,
            task.id_item,
            task.lote,
            cantidad_salida,
            commit=False,
        )
        
        if inventory_result["result"] == 0:
            raise RuntimeError(inventory_result["message"])
        
        # Conservar la tarea como historial y marcarla completada (estado = 3).
        task_update = await picking_repository.update_task_status(
            db, id_task, 3, commit=False
        )
        order_data = await picking_repository.get_order_id_by_picking(db, id_picking)
        if order_data.estado == 1:
            order_data.estado = 2
            order_update = await order_repository.update_order_status(
                db, order_data, commit=False
            )
            if order_update.result == 0:
                raise RuntimeError("Error al actualizar el estado de la orden")
            
        
        
        if task_update["result"] == 0:
            raise RuntimeError("Error al actualizar el estado de la tarea")

        completed = False
        missing_inventory = []
        pending_tasks = await picking_repository.count_pending_picking_tasks(db, id_picking)
        if pending_tasks == 0:
            overview = await get_tasks_overview(db, id_picking)
            missing_inventory = overview.get("missing_inventory", [])
            if not missing_inventory:
                picking_update = await picking_repository.complete_picking(
                    db, id_picking, commit=False
                )
                if picking_update["result"] == 0:
                    raise RuntimeError("Error al completar el picking")
                completed_order = await order_repository.update_order_status_by_picking(
                    db, id_picking, commit=False
                )
                if completed_order.result == 0:
                    raise RuntimeError("Error al completar la orden")
                completed = True

        await db.commit()

        return {
            "result": 1,
            "message": "Tarea confirmada exitosamente",
            "id_movimiento": motion_result.id,
            "completed": completed,
            "missing_inventory": missing_inventory,
        }
        
    except Exception as e:
        print(f"Error al confirmar la tarea: {e}")
        await db.rollback()
        return {"result": 0, "message": f"Error al confirmar la tarea: {str(e)}"}
    
    
async def list_picking_view(db):
    """Lista los pickings con su información detallada para la vista."""
    picking_view_list = await picking_repository.list_picking_view(db)
    return picking_view_list

async def get_list_orders(db):
    return await order_repository.get_list_orders(db)
