import asyncio
import io
import unicodedata
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from time import monotonic

from app.model.aux_inv_model import InvAux, InvAuxCont
from app.repository import motion_repository, inventory_repository, position_repository, item_repository
from app.model.inventory_model import Motion
from app.model.inventory_model import Inventory

POSITIONS_CACHE_TTL_SECONDS = 300

_positions_cache = None
_positions_cache_expires_at = 0.0
_positions_cache_lock = asyncio.Lock()


async def get_positions(db):
    """Obtiene las posiciones y conserva el resultado en memoria durante 5 minutos."""
    global _positions_cache, _positions_cache_expires_at

    now = monotonic()
    if _positions_cache is not None and now < _positions_cache_expires_at:
        return _positions_cache

    async with _positions_cache_lock:
        now = monotonic()
        if _positions_cache is not None and now < _positions_cache_expires_at:
            return _positions_cache

        positions = await inventory_repository.get_positions(db)
        if positions:
            _positions_cache = positions
            _positions_cache_expires_at = monotonic() + POSITIONS_CACHE_TTL_SECONDS

        return positions


def invalidate_positions_cache():
    """Invalida el listado cuando una operación modifica posiciones."""
    global _positions_cache, _positions_cache_expires_at

    _positions_cache = None
    _positions_cache_expires_at = 0.0


async def create_input_invAux(db, invAux:InvAux, invAuxCont:InvAuxCont):
    """Crea el registro de inventario de entrada y manejo de conteo"""
    invAux_data = InvAux()
    invAuxCont_data = InvAuxCont()
    
    #VALIDAR EXISTENCIA DE INVENTARIO AUXILIAR
    existing_invAux = await inventory_repository.get_invAux(db, invAux)
    if existing_invAux.result == 2: # si no existe el registro de inventario auxiliar, se crea uno nuevo
        invAux.estado = 2
        invAux_data = await inventory_repository.create_input_invAux(db, invAux)

        if invAux_data.result != 1:
            return invAux_data

        invAuxCont.id_invAux = invAux_data.id
        invAuxCont.numCont = 1
        invAuxCont.fecha = datetime.now()
        invAuxCont_data = await inventory_repository.create_input_invAuxCont(db, invAuxCont)
        if invAuxCont_data.result == 1:
            invAux_data.result = 1
            invAux_data.message = "Primer conteo registrado exitosamente"
            return {"result": invAux_data.result, "message": invAux_data.message}
        else:
            invAux_data.result = 0
            invAux_data.message = "Error al registrar el primer conteo"
            return {"result": invAux_data.result, "message": invAux_data.message}
    elif existing_invAux.result == 1: # si existe el registro de inventario auxiliar, se crea un nuevo conteo
        invAux_data = existing_invAux
        invAuxCont_data = invAuxCont
        invAuxCont_data.id_invAux = invAux_data.id
        invAuxCont_data.fecha = datetime.now()
        #invAuxCont_data = await inventory_repository.get_invAuxCont(db, invAux_data.id)

        match invAux_data.estado:
            
            case 1: #Pendiente
                invAuxCont_data.numCont = await inventory_repository.get_next_numCont(db, invAux_data.id)
                invAuxCont_data.id_invAux = invAux_data.id
                invAuxCont_data = await inventory_repository.create_input_invAuxCont(db, invAuxCont_data)
                return {"result": invAuxCont_data.result, "message": f"Conteo {invAuxCont_data.numCont} registrado exitosamente"}
                
            case 2: #Conteo 1
                invAuxCont_data.numCont = 2 
                invAux_data.estado = 3 #Conteo 2
                invAuxCont_data.id_invAux = invAux_data.id
                invAuxCont_data = await inventory_repository.create_input_invAuxCont(db, invAuxCont_data)
                invAux_data = await inventory_repository.update_input_invAux(db, invAux_data)
                
                verify = await verify_counts(db,invAux_data)
                return verify
            
            case 4: #Inconsistencia
                invAuxCont_data.numCont = 3
                invAux_data.estado = 1 #Pendiente
                invAuxCont_data.id_invAux = invAux_data.id
                invAuxCont_data = await inventory_repository.create_input_invAuxCont(db, invAuxCont_data)
                invAux_data = await inventory_repository.update_input_invAux(db, invAux_data)
                return {"result": invAuxCont_data.result, "message": "Tercer conteo registrado exitosamente"}
            
            case 6: #Aprobado
                return {"result": 0, "message": "Esa posicion con ese item y lote ya fue aprobada, no se puede registrar un nuevo conteo"}
            
            
async def verify_counts(db,invAux:InvAux):
    """Verifica si los dos conteos coinciden y actualiza el estado del inventario auxiliar en consecuencia."""
    invAux_cont1 = await inventory_repository.get_invAuxCont_by_numCont(db, invAux.id, 1)
    invAux_cont2 = await inventory_repository.get_invAuxCont_by_numCont(db, invAux.id, 2)
    
    if invAux_cont1.result == 1 and invAux_cont2.result == 1:
        if invAux_cont1.cantidad == invAux_cont2.cantidad:
            invAux.estado = 6 #Aprobado
            #Se registra el movimiento de inventario y se actualiza la cantidad en el inventario
            result = await register_inventory_movement(db, invAux, invAux_cont2)
            return result         
        else:
            invAux.estado = 4 #Inconsistencia
            result = 0
            message = f"Inconsistencia detectada: Conteo 1 ({invAux_cont1.cantidad}) y Conteo 2 ({invAux_cont2.cantidad}) no coinciden."
        invAux = await inventory_repository.update_input_invAux(db, invAux)
        return {"result": result, "message": message}

    return {
        "result": 0,
        "message": "No se pueden verificar los conteos porque faltan el conteo 1 y/o el conteo 2."
    }

# Metodo de registro despues de que este aprobado el conteo de esa posicion, se registra el movimiento de inventario y se actualiza la cantidad en el inventario
async def register_inventory_movement(db, invAux:InvAux, invAuxCont:InvAuxCont):
    """Registra el movimiento de inventario y actualiza la cantidad en el inventario."""
    motion_data = Motion()
    inventory_data = Inventory()
    motion_data.tipo_movimiento = 1
    motion_data.id_item = invAux.id_item
    motion_data.lote = invAux.lote
    motion_data.cantidad = invAuxCont.cantidad
    motion_data.posicion_destino_id = invAux.id_posicion
    motion_data.fecha = datetime.now()
    motion_data.id_usuario = invAuxCont.id_usuario
    motion_result = await motion_repository.create_motion_input(db, motion_data)
    if motion_result.result == 1:
        inventory_data.id_item = invAux.id_item
        inventory_data.lote = invAux.lote
        inventory_data.cantidad = invAuxCont.cantidad
        inventory_data.id_posicion = invAux.id_posicion
        inventory_data.fecha_vencimiento = invAux.fecha_vencimiento
        inventory_data.detalles = f"Movimiento de entrada con id {motion_result.id}"
        inventory_result = await inventory_repository.create_input_inventory(db, inventory_data)
        if inventory_result.result == 1:
            await position_repository.update_position_state(db, inventory_data.id_posicion, 2) #Actualizar estado de la posición a ocupada  
            result = 1
            message = "Conteos coinciden. Movimiento de entrada registrado exitosamente."
        else:
            result = 0
            message = "Error al actualizar el inventario." 
    else:
        result = 0
        message = "Error al registrar el movimiento de inventario."
    return {"result": result, "message": message}
    
    
async def list_inv_aux_pending(db):
    """Obtiene el listado de inventario auxiliar pendiente de conteo"""
    inv_aux_list = await inventory_repository.list_inv_aux_pending(db)
    return inv_aux_list


async def get_inventory_consolidated(db):
    """Construye un consolidado por item con total general y desglose por lote."""
    rows = await inventory_repository.get_inventory_consolidated(db)
    consolidated = []
    by_item = {}

    for row in rows:
        item_id = row["id_item"]
        item_entry = by_item.get(item_id)

        if item_entry is None:
            item_entry = {
                "id_item": item_id,
                "cod_item": row["cod_item"],
                "descripcion": row["descripcion"],
                "tipo_item": row["tipo_item"],
                "unidad_medida": row["unidad_medida"],
                "cantidad_total": 0,
                "lotes": [],
            }
            by_item[item_id] = item_entry
            consolidated.append(item_entry)

        cantidad = row["cantidad"] or 0
        item_entry["cantidad_total"] += cantidad
        item_entry["lotes"].append(
            {
                "lote": row["lote"],
                "cantidad": cantidad,
            }
        )

    return consolidated

async def approve_inv_aux(db, id_invAux: int):
    """Aprueba el inventario auxiliar marcándolo directamente con estado 6."""
    invAux = await inventory_repository.get_invAux_by_id(db, id_invAux)
    if invAux.result != 1:
        return {"result": 0, "message": invAux.message}

    invAux.estado = 6
    invAux_updated = await inventory_repository.update_input_invAux(db, invAux)
    if invAux_updated.result != 1:
        return {"result": 0, "message": invAux_updated.message}

    invAuxCont = await inventory_repository.get_invAuxCont(db, invAux_updated.id)
    if invAuxCont.result != 1:
        return {"result": 0, "message": invAuxCont.message}

    result = await register_inventory_movement(db, invAux_updated, invAuxCont)
    if result["result"] != 1:
        return result

    return {"result": 1, "message": "Inventario auxiliar aprobado y registrado exitosamente"}


# ****************************** IMPORTACIÓN MASIVA DE INVENTARIO DESDE EXCEL ******************************

# Alias de encabezados aceptados (normalizados: sin tildes, minúsculas, espacios simples)
IMPORT_COLUMN_ALIASES = {
    "cod_posicion": {"ubicacion", "ubicación", "posicion", "posición", "cod posicion", "cod_posicion"},
    "cod_item": {"referencia", "cod item", "cod_item", "codigo item", "codigo"},
    "lote": {"lote"},
    "fecha_vencimiento": {"f.v", "f v", "fv", "fecha vencimiento", "fecha_vencimiento", "vencimiento"},
    "cantidad": {"cantidad", "cant"},
}

# Columnas opcionales que no se guardan en `inventario` pero se toleran en el archivo
IMPORT_OPTIONAL_COLUMNS = {
    "descripcion": {"descripcion", "descripción"},
}


def _normalise_header(value) -> str:
    value = unicodedata.normalize("NFD", str(value or "").strip().casefold())
    value = "".join(character for character in value if unicodedata.category(character) != "Mn")
    return re.sub(r"\s+", " ", value)


def _normalise_text_value(value) -> str:
    """Convierte un valor de celda a texto, evitando que los números
    (p. ej. referencias o lotes numéricos) queden como '221087.0'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_blank_row(row: dict) -> bool:
    return all(not _normalise_text_value(value) for value in row.values())


def _parse_import_quantity(value):
    raw = _normalise_text_value(value).replace(" ", "")
    if not raw:
        raise ValueError("la cantidad está vacía")
    if "," in raw and "." in raw:
        raw = raw.replace(".", "").replace(",", ".")
    elif "," in raw:
        raw = raw.replace(",", ".")
    try:
        quantity = Decimal(raw)
    except InvalidOperation as error:
        raise ValueError("la cantidad no es un número válido") from error
    if quantity <= 0:
        raise ValueError("la cantidad debe ser mayor que cero")
    return float(quantity)


def _parse_import_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    text_value = _normalise_text_value(value)
    if not text_value:
        return None
    for date_format in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text_value, date_format)
        except ValueError:
            continue
    raise ValueError(f"la fecha de vencimiento '{text_value}' no tiene un formato reconocido")


def _read_inventory_import_rows(content: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows_iterator = workbook.active.iter_rows(values_only=True)
        headers = next(rows_iterator, None)
        if not headers:
            return []
        return [dict(zip(headers, row)) for row in rows_iterator if any(value is not None for value in row)]
    except Exception as error:
        raise ValueError(f"No fue posible leer el archivo XLSX: {error}") from error


def _map_import_columns(rows: list[dict]) -> dict[str, str]:
    if not rows:
        raise ValueError("El archivo no contiene filas de datos")
    headers = {_normalise_header(header): header for header in rows[0].keys()}
    mapping, missing = {}, []
    for field, aliases in IMPORT_COLUMN_ALIASES.items():
        header = next(
            (headers[_normalise_header(alias)] for alias in aliases if _normalise_header(alias) in headers),
            None,
        )
        if header is None:
            missing.append(field)
        else:
            mapping[field] = header
    if missing:
        labels = {
            "cod_posicion": "UBICACIÓN",
            "cod_item": "REFERENCIA",
            "lote": "LOTE",
            "fecha_vencimiento": "F.V",
            "cantidad": "CANTIDAD",
        }
        raise ValueError("Faltan columnas obligatorias: " + ", ".join(labels[item] for item in missing))
    for field, aliases in IMPORT_OPTIONAL_COLUMNS.items():
        mapping[field] = next(
            (headers[_normalise_header(alias)] for alias in aliases if _normalise_header(alias) in headers),
            None,
        )
    return mapping


async def import_inventory_from_excel(db, content: bytes, filename: str = "") -> dict:
    """Actualiza la tabla `inventario` a partir de un archivo Excel de conteo físico.

    Reglas de negocio (confirmadas con el usuario):
    - Si la combinación posición + item + lote ya existe, se SUMA la cantidad del
      archivo a la cantidad existente (no se reemplaza).
    - Si no existe, se crea un nuevo registro de inventario.
    - Los registros que ya están en la BD y no aparecen en el archivo NO se tocan.
    - Las filas cuya posición o referencia no existan en la BD se omiten y se
      reportan al final, sin detener la importación.
    """
    try:
        rows = _read_inventory_import_rows(content)
        columns = _map_import_columns(rows)
    except ValueError as error:
        return {"result": 0, "message": str(error)}

    # Se cargan en memoria los mapas de posiciones e items para no hacer una
    # consulta por cada una de las filas del archivo (puede haber miles).
    positions = await inventory_repository.get_positions(db)
    position_map = {p["cod_posicion"]: p["id"] for p in positions if p.get("cod_posicion")}
    item_map = await item_repository.get_all_items_map(db)

    insertados = 0
    actualizados = 0
    filas_vacias = 0
    errores = []

    for row_number, row in enumerate(rows, start=2):  # fila 1 es el encabezado
        if _is_blank_row(row):
            filas_vacias += 1
            continue

        cod_posicion = _normalise_text_value(row.get(columns["cod_posicion"]))
        cod_item = _normalise_text_value(row.get(columns["cod_item"]))
        lote = _normalise_text_value(row.get(columns["lote"])) or None

        if not cod_posicion:
            errores.append(f"Fila {row_number}: falta la ubicación (posición)")
            continue
        if not cod_item:
            errores.append(f"Fila {row_number}: falta la referencia (item)")
            continue

        id_posicion = position_map.get(cod_posicion)
        if id_posicion is None:
            errores.append(f"Fila {row_number}: la ubicación '{cod_posicion}' no existe en la base de datos")
            continue

        id_item = item_map.get(cod_item)
        if id_item is None:
            errores.append(f"Fila {row_number}: la referencia '{cod_item}' no existe en la base de datos")
            continue

        try:
            cantidad = _parse_import_quantity(row.get(columns["cantidad"]))
        except ValueError as error:
            errores.append(f"Fila {row_number}: {error}")
            continue

        try:
            fecha_vencimiento = _parse_import_date(row.get(columns["fecha_vencimiento"]))
        except ValueError as error:
            errores.append(f"Fila {row_number}: {error}")
            continue

        existente = await inventory_repository.get_inventory_by_position_item_lote(db, id_posicion, id_item, lote)

        if existente:
            resultado = await inventory_repository.increment_inventory_quantity(db, existente["id"], cantidad)
            if resultado.result == 1:
                actualizados += 1
            else:
                errores.append(f"Fila {row_number}: {resultado.message}")
            continue

        nuevo_inventario = Inventory(
            id_posicion=id_posicion,
            id_item=id_item,
            cantidad=cantidad,
            lote=lote,
            fecha_vencimiento=fecha_vencimiento,
            detalles=f"Importado desde archivo Excel: {filename}" if filename else "Importado desde archivo Excel",
        )
        resultado = await inventory_repository.create_input_inventory(db, nuevo_inventario)
        if resultado.result == 1:
            insertados += 1
            await position_repository.update_position_state(db, id_posicion, 2)  # Ocupada
        else:
            errores.append(f"Fila {row_number}: {resultado.message}")

    invalidate_positions_cache()

    total_filas = len(rows)
    procesadas = insertados + actualizados
    mensaje = (
        f"Importación finalizada: {insertados} registros nuevos, {actualizados} actualizados, "
        f"{len(errores)} con error, {filas_vacias} filas vacías omitidas de {total_filas} filas leídas."
    )

    return {
        "result": 1 if procesadas > 0 or not errores else 0,
        "message": mensaje,
        "total_filas": total_filas,
        "insertados": insertados,
        "actualizados": actualizados,
        "filas_vacias": filas_vacias,
        "errores": errores,
    }